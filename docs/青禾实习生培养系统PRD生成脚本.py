# -*- coding: utf-8 -*-
"""生成《青禾·实习生培养系统 PRD v1.0》正式文档"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Microsoft YaHei'
BRAND = RGBColor(0x14, 0x60, 0x3F)
DARK = RGBColor(0x1C, 0x2B, 0x23)
GRAY = RGBColor(0x5B, 0x6D, 0x62)

def _set_rFonts(el, name):
    rPr = el.get_or_add_rPr()
    rF = rPr.get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rF.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        k = qn(a)
        if rF.get(k) is not None:
            del rF.attrib[k]

def set_doc_default_font(doc, name=FONT):
    styles_el = doc.styles.element
    dd = styles_el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults'); styles_el.insert(0, dd)
    rpd = dd.find(qn('w:rPrDefault'))
    if rpd is None:
        rpd = OxmlElement('w:rPrDefault'); dd.append(rpd)
    rPr = rpd.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr'); rpd.append(rPr)
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts'); rPr.insert(0, rF)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rF.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        k = qn(a)
        if rF.get(k) is not None:
            del rF.attrib[k]

def apply_chinese_fonts(doc, name=FONT):
    set_doc_default_font(doc, name)
    _set_rFonts(doc.styles['Normal'].element, name)
    for i in range(1, 10):
        try: _set_rFonts(doc.styles[f'Heading {i}'].element, name)
        except KeyError: pass
    try: _set_rFonts(doc.styles['Title'].element, name)
    except KeyError: pass

def set_run(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    _set_rFonts(run._element, name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = color

def para(doc, text, size=10.5, bold=False, color=None, space_after=6, indent=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    if align: p.alignment = align
    return p

def heading(doc, text, level, color=DARK):
    h = doc.add_heading('', level=level)
    r = h.add_run(text)
    set_run(r, size={1:16, 2:13.5, 3:12}.get(level, 11), bold=True, color=BRAND if level == 1 else (BRAND if level == 2 else DARK))
    h.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
    h.paragraph_format.space_after = Pt(6)
    return h

def bullet(doc, text, level=0, size=10.5):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    r = p.add_run(text)
    set_run(r, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def make_table(doc, headers, rows, widths=None, header_fill='E5F3EB'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(h)
        set_run(r, size=9.5, bold=True, color=BRAND)
        shade_cell(hdr[i], header_fill)
        if widths: hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            set_run(r, size=9)
            if widths: cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run('【说明】' + text)
    set_run(r, size=9, color=GRAY)
    p.paragraph_format.space_after = Pt(6)
    return p

doc = Document()
# 页边距 A4
for sec in doc.sections:
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
apply_chinese_fonts(doc)

# ============ 封面 ============
for _ in range(5): doc.add_paragraph()
p = para(doc, '青禾 · 数智化服务平台', size=13, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
p = para(doc, '实习生培养系统', size=30, bold=True, color=BRAND, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
p = para(doc, '产品需求文档（PRD）', size=16, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p = para(doc, '双岗位模型 · 三角色视角 · 全周期培养闭环', size=11, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
doc.add_paragraph()
make_table(doc,
    ['项目', '内容'],
    [
        ['文档版本', 'V1.0'],
        ['文档状态', '待评审'],
        ['产品名称', '青禾 · 实习生培养系统（GROW TOGETHER）'],
        ['编写人', '产品经理'],
        ['日期', '2026-09-06'],
        ['关联产物', '高保真交互原型 index.html（同目录）'],
        ['适用范围', '数智化服务平台 · 技术团队与 AI 产品团队实习培养'],
    ],
    widths=[4.0, 12.0])
doc.add_page_break()

# ============ 1 产品概述 ============
heading(doc, '1. 产品概述', 1)
heading(doc, '1.1 背景与问题', 2)
para(doc, '公司每年接收大量开发/产品实习生，但培养过程存在三个核心问题：')
bullet(doc, '过程黑盒：实习生每周在做什么、进展如何、遇到什么卡点，管理者缺乏统一视图；')
bullet(doc, '培养不标准：各导师带教方式差异大，无统一的阶段化目标与验收机制；')
bullet(doc, '转正靠印象：转正评估缺少过程数据支撑，评审材料临时拼凑，结论难服众。')
heading(doc, '1.2 产品目标', 2)
make_table(doc, ['目标', '说明', '衡量方式'],
    [
        ['让成长看得见', '周粒度呈现每位实习生的目标、任务、进步与卡点', '每周培养总览使用率 ≥ 80%'],
        ['让培养有节奏', '按岗位定义 4 阶段成长路径与晋升规则', '阶段晋升均有过程证据'],
        ['让转正有依据', '转正评估基于沉淀的过程数据与能力评分', '转正材料完整率 100%'],
    ],
    widths=[3.2, 7.4, 5.4])
heading(doc, '1.3 设计原则', 2)
bullet(doc, '按入职阶段关注，不做排名：避免横向竞争内卷，培养只与阶段标准对齐；')
bullet(doc, '双岗位模型：开发与产品各自独立的阶段、能力维度、评估体系，系统层面统一；')
bullet(doc, '过程数据自动沉淀：任务、1:1、阶段回顾、能力自评自动汇入成长档案；')
bullet(doc, '角色闭环：导师管培养、实习生有动作、HR 看组织，三方各司其职。')

# ============ 2 用户角色与权限 ============
heading(doc, '2. 用户角色与权限', 1)
heading(doc, '2.1 角色定义', 2)
make_table(doc, ['角色', '定位', '核心职责'],
    [
        ['带教导师（技术导师 / 产品导师）', '培养执行者', '制定周目标、下发/验收任务、记录 1:1、协调卡点、确认阶段晋升；导师视角标签随岗位模式联动（开发模式=技术导师·技术视角，产品模式=产品导师·产品视角）'],
        ['实习生', '被培养者', '接收任务、提交成果、阶段自评、提报卡点、准备转正材料'],
        ['HR / 组织管理员', '组织管理者', '跨岗位组织看板、转正流程监控、岗位模型配置、数据导出'],
    ],
    widths=[4.4, 3.0, 8.6])
heading(doc, '2.2 权限矩阵', 2)
make_table(doc, ['功能', '带教导师', '实习生', 'HR 管理员'],
    [
        ['培养总览（周视图）', '本人名下', '—', '—'],
        ['个人成长档案', '本人名下', '本人', '全部'],
        ['创建成长任务', '✓', '—', '—'],
        ['提交成果 / 阶段自评', '—', '✓', '—'],
        ['验收任务', '本人名下', '—', '—'],
        ['记录 / 查看 1:1', '本人名下', '本人', '全部'],
        ['提报卡点（支持请求）', '✓', '✓', '—'],
        ['协调支持', '✓', '—', '✓'],
        ['阶段晋升确认', '✓', '—', '✓ 复核'],
        ['转正评估', '名下实习生', '—', '全流程'],
        ['组织看板 / 导师负载', '—', '—', '✓'],
        ['岗位模型配置（阶段/权重）', '—', '—', '✓'],
        ['实习生管理（添加/编辑/删除）', '✓', '—', '✓'],
        ['重置示例数据', '✓（DEMO）', '—', '—'],
    ],
    widths=[6.2, 3.4, 3.4, 3.0])

# ============ 3 核心概念 ============
heading(doc, '3. 核心概念与术语', 1)
make_table(doc, ['术语', '定义'],
    [
        ['培养阶段', '按岗位定义的成长路径（开发/产品各 4 阶段），实习生按阶段对齐期望'],
        ['阶段进度', 'x/4 表示当前所处阶段与总阶段数，由晋升规则驱动前进'],
        ['成长任务', '服务本周成长目标的具象任务，有标题、验收标准、状态流转'],
        ['周成长目标', '每周一个可验收的一句话目标，与任务联动'],
        ['1:1 记录', '导师与实习生的固定沟通，含议题、结论、行动项'],
        ['支持请求（卡点）', '实习生或导师提报的阻塞项，待协调→已协调'],
        ['进步记录', '里程碑式进步，可打星标，汇入成长档案'],
        ['能力模型', '按岗位定义的 5 维能力（如产品：需求洞察/方案设计/数据意识/文档表达/协作推进）'],
        ['阶段自评', '实习生对照能力维度自评（1-5 分），导师确认后生效'],
        ['转正评估', '5 维 × 权重 = 综合分，基于过程数据沉淀评分证据'],
        ['转正考核版本', '产品岗转正考核体系版本：V1 维度权重评估（原）/ V2 三维度必要条件（最新），两版内容均保留可查'],
        ['必要条件（V2）', '产品岗转正必备：①子领域全流程独立承担 ≥1 ②AI 场景设计-落地-演进 ≥2 ③量化产出（5 需求文档/5 需求澄清/2 竞品分析/2 数据运营）'],
        ['量化产出清单', 'V2 考核下的可量化成果达成度（x/目标），未达标项警示'],
        ['加分项（Skill）', '沉淀 Skill 给团队使用（Agent 技能/提效模板等），转正评估加分'],
    ],
    widths=[3.4, 12.6])

# ============ 4 信息架构 ============
heading(doc, '4. 信息架构与导航', 1)
para(doc, '系统按角色提供三套导航结构：', space_after=4)
make_table(doc, ['角色', '导航菜单'],
    [
        ['带教导师 TL', '培养总览 / 个人成长 / 成长任务(待验收角标) / 沟通与支持 / 阶段回顾 / 转正准备'],
        ['实习生', '我的成长 / 我的任务(待验收角标) / 我的沟通 / 转正申请'],
        ['HR 管理员', '组织看板 / 成长档案 / 转正流程(窗口内人数角标)'],
    ],
    widths=[3.4, 12.6])
para(doc, '全局控件：顶栏含岗位模式切换（开发实习生/产品实习生）、角色视角切换（TL/实习生/HR）、DEMO 标识。', space_after=6)

# ============ 5 功能需求 ============
heading(doc, '5. 功能需求详述', 1)

heading(doc, '5.1 岗位模式切换（全局）', 2)
para(doc, '功能描述：顶栏分段控件，一键在「开发实习生 / 产品实习生」两套培养模型间切换。', space_after=4)
make_table(doc, ['编号', '需求点', '验收要点'],
    [
        ['F1.1', '切换岗位模式', '切换后全站培养模型联动替换：阶段定义、能力维度、转正评估权重、示例数据'],
        ['F1.2', '状态保留', '切换岗位时重置示例数据；视角（TL/实习生/HR）保持不变'],
        ['F1.3', '数据隔离', '两岗位数据完全隔离，互不串扰'],
    ],
    widths=[1.6, 5.0, 9.4])

heading(doc, '5.2 培养总览（F2）', 2)
para(doc, '功能描述：导师视角核心驾驶舱（产品模式=产品导师视角，开发模式=技术导师视角），周粒度呈现名下实习生整体状态。', space_after=4)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F2.1', '头部品牌区', '标识 QINGHE/GROWTH WORKSPACE + 主标题「每个人的成长，都看得见。」+ 副标题'],
        ['F2.2', '本周视图条', '本周日期范围 + 统计：伙伴数 / 待验收数 / 需要支持数，数字可点击筛选卡片'],
        ['F2.3', '实习生卡片流', '每人一卡：头像/姓名/岗位周数/阶段标签(x/4)/阶段进度条/本周目标/当前任务+状态/最近进步⭐/需要支持块/下次1:1/查看成长档案'],
        ['F2.4', '新建成长任务', '弹窗表单：标题*、指派伙伴*、说明；创建后任务进入该伙伴列表（进行中）'],
        ['F2.5', '任务验收入口', '待验收任务卡片提供「去验收」按钮，进入验收弹窗'],
        ['F2.6', '卡点协调入口', '待协调支持块提供「协调支持」按钮，点击后状态置为已协调'],
        ['F2.7', '筛选与空态', '统计点击筛选后显示筛选提示与清除入口；无结果显示空态'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.3 个人成长（F3）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F3.1', '档案头部卡', '头像/姓名/岗位周数/入职日期/带教导师/阶段标签与进度条/当前周目标'],
        ['F3.2', '能力维度', '按岗位 5 维能力横向条形展示（当前分/5），实习生视角可发起「阶段自评」与「提报卡点」'],
        ['F3.3', '成长轨迹', '按周时间线：每周里程碑事项，当前周高亮标记'],
        ['F3.4', '1:1 记录', '历史 1:1 列表（日期/议题/结论/行动项）'],
        ['F3.5', '进步记录', '里程碑进步列表，星标项突出展示'],
        ['F3.6', '档案切换（TL）', 'TL 可切换查看名下任一实习生档案'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.4 成长任务（F4）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F4.1', '任务列表', '全部任务平铺，含伙伴、带教、状态标签'],
        ['F4.2', '状态筛选 Tab', '全部/进行中/待验收/已验收，各 Tab 显示数量'],
        ['F4.3', '状态流转', '进行中→待验收（实习生「提交成果」或导师「提交验收」）→已验收（导师「去验收」）'],
        ['F4.4', '提交成果（实习生）', '弹窗：成果说明*、产出物链接；提交后进入待验收'],
        ['F4.5', '验收（导师）', '弹窗：验收结论；确认后置为已验收，并沉淀为进步记录'],
        ['F4.6', '新建任务（导师）', '同 F2.4，入口复用'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.5 沟通与支持（F5）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F5.1', '1:1 安排', '展示名下次 1:1（日期/带教/时长建议），已排期状态'],
        ['F5.2', '支持请求', '卡点列表：发起人、描述、状态（待协调/已协调）；TL/HR 可执行协调'],
        ['F5.3', '实习生视图', '仅展示本人的 1:1 与支持请求'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.6 阶段回顾（F6）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F6.1', '阶段卡片', '按岗位 4 阶段：阶段名/周期/描述/能力标签/状态（已完成/进行中/未开始）'],
        ['F6.2', '晋升规则', '3 项规则卡片：任务验收≥2、自评+导师确认、关键能力≥3/5；满足全部自动晋升'],
        ['F6.3', '回顾详情', '最近一次阶段回顾：做得好的/待提升的/下一步行动/导师评价'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.7 转正准备（F7）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F7.1', '流程步骤条', '材料准备→必要条件核验（V2）→导师评估→评审会→结果归档'],
        ['F7.2', '考核版本切换（产品岗）', 'V1 维度权重评估 / V2 三维度必要条件，顶栏分段控件一键切换；两版内容均保留可查'],
        ['F7.3', 'V1 评估维度', '按岗位 5 维 × 权重展示，含锚点说明（1-5 分制，3 分为合格线）'],
        ['F7.4', 'V2 三维度必要条件', '①能力·子领域全流程：独立承担某子领域「规划→需求拆解/细化/优先级→需求澄清落地→项目跟进→上线验收→数据运营」全流程；②实践·AI 场景：最少主导参与 2 个 AI 场景的设计-落地-演进；③综合·量化产出：5 份需求文档 / 5 场需求澄清 / 2 份竞品分析 / 2 场数据运营（能通过数据识别问题并提出方案）'],
        ['F7.5', '量化产出清单', '逐项展示达成进度（x/目标），未达标项琥珀色警示；全部达成 + 导师评估通过方可进入评审'],
        ['F7.6', '加分项', '沉淀 Skill 给团队使用（Agent 技能/提效模板/自动化脚本）为加分项，原型含已加分示例'],
        ['F7.7', '伙伴切换', '导师/HR 视角可按伙伴查看各自量化进度'],
        ['F7.8', '材料模板', '提供评估模板下载'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.8 组织看板（F8，HR）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F8.1', '统计大卡', '在培实习生总数（含岗位分布）/待验收/需要支持/转正窗口内人数'],
        ['F8.2', '培养阶段分布', '跨岗位阶段分布色条 + 图例；岗位分布统计'],
        ['F8.3', '带教导师一览', '每位导师：带教人数（按岗位拆分）、待验收数、需支持数；负载建议 ≤3 人'],
        ['F8.4', '转正流程监控', '全流程节点状态 + 当前窗口人数提示'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.9 系统级功能（F9）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F9.1', '角色视角切换', '顶栏下拉：带教导师（技术导师/产品导师，随岗位联动）/ 实习生 / 组织管理 HR，菜单与内容联动'],
        ['F9.2', '导师视角岗位联动', '岗位模式切换为开发时，导师视角标签显示「技术导师·技术视角」，头像标记「技」；切换为产品时显示「产品导师·产品视角」，头像标记「产」'],
        ['F9.3', '重置示例数据（DEMO）', '一键恢复示例数据，用于演示'],
        ['F9.4', '导航角标', '成长任务（导师/实习生）显示待验收数；转正流程（HR）显示窗口内人数'],
    ],
    widths=[1.6, 4.6, 9.8])

heading(doc, '5.10 实习生管理（F10）', 2)
make_table(doc, ['编号', '需求点', '说明'],
    [
        ['F10.1', '列表展示', '当前岗位模式下全部实习生列表：姓名/岗位/周数/阶段/带教/入职/本周目标'],
        ['F10.2', '添加实习生', '表单：姓名*、岗位*、入职周数*、培养阶段、带教导师*、本周成长目标*；新伙伴自动获得默认培养数据（能力评分、量化产出、成长轨迹起始项）并出现在培养总览/组织看板'],
        ['F10.3', '编辑实习生', '修改信息实时同步至成长档案、培养总览、转正页'],
        ['F10.4', '删除实习生', '二次确认弹窗；删除后该实习生的任务/1:1/回顾等记录一并移除；删除后索引与统计自动修正'],
        ['F10.5', '岗位联动', '开发/产品模式各自管理本岗位实习生；组织看板统计实时联动（当前岗位取实时数据，另一岗位取示例数据）'],
        ['F10.6', '权限', '带教导师与 HR 可管理；实习生视角不提供管理入口'],
    ],
    widths=[1.6, 4.6, 9.8])

# ============ 6 数据模型 ============
heading(doc, '6. 数据模型（字段级）', 1)
heading(doc, '6.1 实习生 Intern', 2)
make_table(doc, ['字段', '类型', '必填', '说明'],
    [
        ['id', 'string', '是', '唯一标识'],
        ['name', 'string', '是', '姓名'],
        ['role', 'string', '是', '岗位（开发：后端开发/质量工程/AI应用开发；产品：AI 产品）'],
        ['jobType', 'enum', '是', '岗位大类：dev / product'],
        ['week', 'int', '是', '入职第几周'],
        ['phase', 'int', '是', '当前培养阶段索引 0-3'],
        ['goal', 'string', '是', '本周成长目标'],
        ['task', 'object', '是', '当前任务 {title, status}'],
        ['progress', 'object', '否', '最近进步 {text, star}'],
        ['support', 'object', '否', '支持请求 {text, status}'],
        ['meeting', 'string', '是', '下次 1:1 日期'],
        ['tl', 'string', '是', '带教导师'],
        ['skills', 'map', '是', '5 维能力评分（key→1-5）'],
        ['quant', 'map', '否', '转正量化产出（V2，产品岗）：{doc 需求文档数, clarify 需求澄清场次, cmp 竞品分析数, ops 数据运营场次, ai AI场景数, flow 子领域闭环数, skill 是否沉淀加分项}'],
        ['timeline', 'array', '是', '成长轨迹 [{w, t, d, cur}]'],
        ['m1', 'array', '是', '1:1 记录 [{d, c}]'],
        ['progressLog', 'array', '是', '进步记录 [{t, d, star}]'],
        ['join', 'date', '是', '入职日期'],
    ],
    widths=[3.0, 2.2, 1.4, 9.4])
heading(doc, '6.2 成长任务 Task', 2)
make_table(doc, ['字段', '类型', '必填', '说明'],
    [
        ['id', 'string', '是', '任务唯一标识'],
        ['internId', 'string', '是', '归属实习生'],
        ['title', 'string', '是', '任务标题'],
        ['status', 'enum', '是', 'doing 进行中 / review 待验收 / done 已验收'],
        ['desc', 'string', '否', '任务说明与验收标准'],
        ['submitNote', 'string', '否', '实习生提交成果说明'],
        ['submitLink', 'string', '否', '产出物链接'],
        ['reviewNote', 'string', '否', '导师验收结论'],
        ['createdAt / updatedAt', 'datetime', '是', '创建/更新时间'],
    ],
    widths=[3.0, 2.2, 1.4, 9.4])
heading(doc, '6.3 培养阶段 Stage', 2)
make_table(doc, ['字段', '类型', '必填', '说明'],
    [
        ['jobType', 'enum', '是', 'dev / product'],
        ['index', 'int', '是', '阶段序号 0-3'],
        ['name', 'string', '是', '阶段名（开发：熟悉融入→跟随式交付→独立承担→转正准备；产品：跟学理解→需求分析→独立设计与验证→转正准备）'],
        ['range', 'string', '是', '建议周期（如第 1-3 周）'],
        ['desc', 'string', '是', '阶段目标描述'],
        ['skills', 'array', '是', '本阶段能力标签'],
        ['promoteRules', 'array', '是', '晋升规则（验收≥2 / 自评+导师确认 / 关键能力≥3分）'],
    ],
    widths=[3.0, 2.2, 1.4, 9.4])

# ============ 7 状态机 ============
heading(doc, '7. 状态机定义', 1)
heading(doc, '7.1 成长任务状态机', 2)
make_table(doc, ['当前状态', '触发事件', '执行角色', '目标状态'],
    [
        ['进行中 doing', '实习生提交成果 / 导师代提交验收', '实习生 / TL', '待验收 review'],
        ['待验收 review', '导师验收通过', 'TL', '已验收 done'],
        ['待验收 review', '验收驳回（原型未实现，预留）', 'TL', '进行中 doing'],
        ['已验收 done', '—（终态）', '—', '—'],
    ],
    widths=[3.6, 5.6, 2.6, 3.4])
heading(doc, '7.2 支持请求状态机', 2)
make_table(doc, ['当前状态', '触发事件', '执行角色', '目标状态'],
    [
        ['待协调 pending', 'TL / HR 执行协调', 'TL / HR', '已协调 resolved'],
    ],
    widths=[3.6, 5.6, 2.6, 3.4])
heading(doc, '7.3 培养阶段状态机（晋升）', 2)
para(doc, '阶段状态：未开始 → 进行中 → 已完成。晋升条件（满足全部 3 项，自动进入下一阶段）：', space_after=4)
bullet(doc, '本阶段成长任务验收 ≥ 2 个，且无未闭环事项；')
bullet(doc, '实习生完成阶段自评，导师 1:1 确认成长证据；')
bullet(doc, '本阶段核心能力项导师评分 ≥ 3/5（5 分制）。')
heading(doc, '7.4 转正流程状态机', 2)
make_table(doc, ['节点', '输入', '责任角色', '输出'],
    [
        ['材料准备', '成长档案 + 成果清单（自动沉淀）', '实习生', '材料包'],
        ['必要条件核验（V2）', '三维度必要条件 + 量化产出清单，系统自动核验缺项', '系统 / 实习生', '核验结果'],
        ['导师评估', '导师按考核版本打分并填写评语', '产品导师', '评估表'],
        ['评审会', '答辩 + 集体评议', 'HR 组织 / 评审组', '评审结论'],
        ['结果归档', 'HR 确认并归档', 'HR', '转正结果 + 反馈'],
    ],
    widths=[3.0, 6.2, 3.0, 3.6])

# ============ 8 业务流程 ============
heading(doc, '8. 关键业务流程', 1)
heading(doc, '8.1 周培养闭环', 2)
para(doc, '每周循环：导师设定周目标 → 下发成长任务 → 实习生执行并提交成果 → 导师验收 → 1:1 复盘（确认进步/卡点/下周目标）→ 数据沉淀入档案。培养总览页即该循环的驾驶舱。', space_after=6)
heading(doc, '8.2 阶段晋升流程', 2)
para(doc, '实习生完成阶段任务 → 发起阶段自评 → 导师 1:1 确认 → 系统校验 3 项晋升规则 → 满足则阶段 +1，不满足则给出缺口提示（原型以规则卡片展示，正式版需缺口明细）。', space_after=6)
heading(doc, '8.3 转正评估链', 2)
para(doc, '进入最后培养阶段 → 自动开放转正流程 → 过程数据（任务/1:1/回顾/自评）自动汇入评估材料 → 导师评估 → 评审会 → 结果归档。评估全程留痕，结论可追溯。', space_after=6)

# ============ 9 双岗位模型 ============
heading(doc, '9. 双岗位模型配置（可配置化）', 1)
heading(doc, '9.1 培养阶段对比', 2)
make_table(doc, ['阶段', '开发实习生', '产品实习生'],
    [
        ['阶段一', '熟悉融入（第 1-2 周）：了解业务、代码库与协作规范', '跟学理解（第 1-3 周）：理解产品、用户与需求流转流程'],
        ['阶段二', '跟随式交付（第 3-6 周）：指导下完成真实任务，学会复现与复盘', '需求分析（第 4-7 周）：独立用户调研与需求洞察，输出需求文档'],
        ['阶段三', '独立承担（第 7-10 周）：独立负责模块，用证据验证效果', '独立设计与验证（第 8-11 周）：独立模块设计，用数据验证效果'],
        ['阶段四', '转正准备（第 11-12 周）：沉淀成果与材料', '转正准备（第 12 周）：沉淀成果与材料'],
    ],
    widths=[2.0, 7.0, 7.0])
heading(doc, '9.2 能力模型与转正评估权重', 2)
make_table(doc, ['维度', '开发权重', '产品权重'],
    [
        ['开发：工程实现 30% / 问题定位 25% / 交付质量 20% / 技术沟通 15% / 学习成长 10%', '—', '—'],
        ['产品：需求洞察 25% / 方案设计 25% / 数据意识 20% / 文档表达 15% / 协作推进 15%', '—', '—'],
    ],
    widths=[8.0, 4.0, 4.0])
note(doc, '权重与阶段周期需在配置后台开放编辑，供 HR 按组织制度调整；原型中为示例值。')
heading(doc, '9.3 转正考核版本（产品岗）', 2)
make_table(doc, ['版本', '体系', '说明'],
    [
        ['V1', '维度权重评估', '5 维 × 权重（需求洞察 25% / 方案设计 25% / 数据意识 20% / 文档表达 15% / 协作推进 15%），1-5 分制加权'],
        ['V2（最新）', '三维度必要条件', '①能力·子领域全流程独立承担 ≥1 ②AI 场景设计-落地-演进 ≥2 ③量化产出：5 需求文档 / 5 需求澄清 / 2 竞品分析 / 2 数据运营（可数据识别问题+方案）；加分项：沉淀 Skill 给团队'],
        ['切换规则', '版本可一键切换，两版内容均保留', '切换不丢失任何一方内容；V2 未达标项在量化清单中警示'],
    ],
    widths=[3.0, 6.6, 6.4])
note(doc, '需求澄清场次数示例取 5 场（与 5 份需求文档对齐），阈值需在配置后台开放编辑。')

# ============ 10 非功能需求 ============
heading(doc, '10. 非功能需求', 1)
make_table(doc, ['类别', '要求'],
    [
        ['权限与安全', '基于角色的访问控制（RBAC）；实习生仅可见本人数据，导师仅可见名下数据；敏感字段脱敏'],
        ['数据与审计', '所有培养动作（任务/验收/1:1/自评/晋升）留痕，支持追溯'],
        ['性能', '培养总览首屏加载 ≤ 2s；支持 50 人团队规模流畅操作'],
        ['兼容性', '桌面端 Chrome / Edge 最新两个大版本；移动端仅支持查看'],
        ['可配置性', '阶段定义、晋升阈值、评估权重、岗位模型均配置化，无需发版'],
        ['可扩展性', '岗位模型可扩展（如设计实习生/运营实习生），支持新增岗位类型'],
        ['数据对接（预留）', '开发岗：代码仓库/缺陷系统；产品岗：需求池/数据看板；初期支持手动录入与模板导入'],
    ],
    widths=[3.4, 12.6])

# ============ 11 埋点与度量 ============
heading(doc, '11. 埋点与度量指标', 1)
heading(doc, '11.1 北极星指标', 2)
para(doc, '北极星：转正决策有过程数据支撑的比例（目标 100%——每一位转正实习生的评估材料均自动沉淀过程证据）。', space_after=6)
heading(doc, '11.2 指标体系', 2)
make_table(doc, ['层级', '指标', '口径'],
    [
        ['过程', '周任务验收率', '本周任务验收数 / 本周任务下发数'],
        ['过程', '1:1 完成率', '实际完成 1:1 数 / 计划数'],
        ['过程', '卡点平均解决时长', '提报→已协调的平均天数'],
        ['过程', '阶段晋升平均周数', '跨阶段所需周数（对照建议周期）'],
        ['结果', '转正通过率', '评审通过 / 进入流程数'],
        ['结果', '评估材料完整率', '含任务+1:1+回顾+自评的材料占比'],
        ['结果', '导师带教满意度', '实习生季度反馈评分'],
    ],
    widths=[2.4, 5.0, 8.6])
heading(doc, '11.3 关键埋点事件', 2)
bullet(doc, '任务：create_task / submit_work / review_pass / review_reject；')
bullet(doc, '沟通：m1_done / support_raise / support_resolve；')
bullet(doc, '成长：self_review_submit / phase_up / transfer_flow_start / transfer_result / check_progress；')
bullet(doc, '使用：view_overview / view_archive / switch_role / switch_job / switch_check_ver。')

# ============ 12 验收标准 ============
heading(doc, '12. 验收标准（MVP）', 1)
make_table(doc, ['#', '验收项', '验收标准'],
    [
        ['A1', '岗位切换', '开发/产品两套模型一键切换，阶段/能力/评估/数据全联动'],
        ['A2', '培养总览', '统计数字、卡片字段、筛选、空态全部正确'],
        ['A3', '任务闭环', '进行中→待验收→已验收全链路，统计与角标联动'],
        ['A4', '实习生动作', '提交成果、阶段自评、提报卡点可用，状态正确流转'],
        ['A5', 'HR 组织看板', '跨岗位统计、导师负载、转正流程监控正确'],
        ['A6', '权限隔离', '实习生无法看到他人档案与任务'],
        ['A7', '数据重置（DEMO）', '一键恢复初始示例数据'],
        ['A8', '权限矩阵', '按 2.2 节矩阵逐项核对通过'],
        ['A9', '实习生管理', '添加/编辑/删除可用，二次确认生效，数据实时联动总览统计与组织看板'],
        ['A10', '删除边界', '删除后列表索引、个人成长选中态、统计数字自动修正，无残留引用'],
    ],
    widths=[1.4, 4.0, 10.6])

# ============ 13 里程碑 ============
heading(doc, '13. 里程碑与排期', 1)
make_table(doc, ['阶段', '范围', '周期', '交付物'],
    [
        ['P0 MVP', '岗位切换 + 培养总览 + 成长任务 + 个人成长 + 实习生管理', '4 周', '可用的最小培养闭环'],
        ['P1', '沟通与支持 + 阶段回顾 + HR 组织看板', '+3 周', '三角色闭环'],
        ['P2', '转正流程 + 权限体系 + 数据对接 + 配置后台', '+2 周', '完整系统上线'],
    ],
    widths=[2.4, 8.0, 2.0, 3.6])

# ============ 14 待确认事项 ============
heading(doc, '14. 待确认事项（Open Questions）', 1)
make_table(doc, ['#', '事项', '影响'],
    [
        ['Q1', '阶段周期与晋升阈值是否与公司现行制度一致（示例：验收≥2、评分≥3/5）', '晋升规则参数'],
        ['Q2', '转正评审线上化程度：仅材料准备，还是含评审会审批流与电子签', 'P2 范围'],
        ['Q3', '数据对接源：开发岗接代码仓库/缺陷系统？产品岗接需求池/数据平台？', 'P2 集成设计'],
        ['Q4', '通知渠道：飞书 / 企业微信 / 邮件（验收、1:1 提醒、卡点升级）', 'P1 范围'],
        ['Q5', '岗位模型是否还需支持第三类岗位（如设计/运营实习生）', '架构预留'],
        ['Q6', '「需要支持」的升级机制：超过 N 天未协调是否自动升级提醒', 'F5 增强'],
        ['Q7', 'V2 量化阈值确认：需求澄清场次示例为 5 场、需求文档 5 份、竞品 2 份、数据运营 2 场、AI 场景 2 个——是否与团队现行标准一致，阈值配置化', 'V2 考核参数'],
    ],
    widths=[1.4, 10.6, 4.0])

# ============ 页脚 ============
sec = doc.sections[0]
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer_p.add_run('青禾 · 实习生培养系统 PRD V1.0  |  机密 · 仅限内部使用')
set_run(fr, size=8, color=GRAY)

out = r'C:\Users\Sangfor\Documents\Loomy Workspace\青禾实习生培养系统PRD.docx'
doc.save(out)
print('SAVED:', out)
